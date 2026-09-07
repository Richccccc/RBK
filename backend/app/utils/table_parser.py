"""表格文件解析：xlsx（openpyxl）与 csv（标准库），统一转成 {name, headers, rows} 结构。

- 单元格一律转字符串；日期保留可读格式；空单元格为空串
- 限制单 sheet 行数，防止异常大文件拖垮免费服务器
"""
import csv
import io
from datetime import date, datetime, time
from typing import Any

MAX_ROWS_PER_SHEET = 20000  # 每个 sheet 最多解析行数（含表头），防异常大文件拖垮服务器
MAX_COLS = 64  # 每行最多列数


class SheetParseError(Exception):
    """解析失败。"""


def _cell_str(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        if v.time() == time(0, 0):
            return v.strftime("%Y-%m-%d")
        return v.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(v, date):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, time):
        return v.strftime("%H:%M:%S")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def parse_xlsx(data: bytes) -> list[dict]:
    try:
        from openpyxl import load_workbook
    except ImportError as e:  # pragma: no cover
        raise SheetParseError("服务器缺少 openpyxl，无法解析 xlsx") from e

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:
        raise SheetParseError("无法读取 xlsx 文件，文件可能已损坏") from e

    tables: list[dict] = []
    try:
        for ws in wb.worksheets:
            rows_raw: list[list[str]] = []
            for row in ws.iter_rows(max_row=MAX_ROWS_PER_SHEET, values_only=True):
                rows_raw.append([_cell_str(v) for v in row[:MAX_COLS]])

            # 剔除尾部全空行
            while rows_raw and not any(rows_raw[-1]):
                rows_raw.pop()
            # 剔除尾部全空列（read_only 模式可能按标称维度吐出空列）
            width = 0
            for r in rows_raw:
                for j in range(len(r) - 1, -1, -1):
                    if r[j]:
                        width = max(width, j + 1)
                        break
            rows_raw = [r[:width] for r in rows_raw]

            table = {"name": ws.title or "Sheet", "headers": [], "rows": []}
            if not rows_raw:
                tables.append(table)
                continue

            # 首行作表头；全空首行则用 列1/列2… 占位
            first = rows_raw[0]
            if any(first):
                table["headers"] = [c or f"列{j + 1}" for j, c in enumerate(first)]
                table["rows"] = rows_raw[1:]
            else:
                table["headers"] = [f"列{j + 1}" for j in range(width)]
                table["rows"] = rows_raw
            tables.append(table)
    finally:
        wb.close()

    if not tables:
        raise SheetParseError("文件中没有工作表")
    return tables


def parse_csv(data: bytes) -> list[dict]:
    # 尝试常见编码，utf-8 优先（兼容 Excel 导出的 GBK csv）
    text = None
    for enc in ("utf-8-sig", "utf-8", "gbk", "gb18030"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise SheetParseError("无法识别 CSV 文件编码（支持 UTF-8 / GBK）")

    reader = csv.reader(io.StringIO(text))
    raw = [[_cell_str(c) for c in row[:MAX_COLS]] for _, row in zip(range(MAX_ROWS_PER_SHEET), reader)]
    if not raw:
        raise SheetParseError("CSV 文件为空")

    headers = raw[0]
    if not any(headers):
        headers = [f"列{j + 1}" for j in range(len(headers))]
    rows = raw[1:]
    return [{"name": "CSV", "headers": [h or f"列{j + 1}" for j, h in enumerate(headers)], "rows": rows}]


def parse_table_file(data: bytes, filename: str) -> tuple[list[dict], int]:
    """按扩展名解析，返回 (tables, 总数据行数)。不支持的格式抛 SheetParseError。"""
    lower = (filename or "").lower()
    if lower.endswith(".xlsx"):
        tables = parse_xlsx(data)
    elif lower.endswith(".csv"):
        tables = parse_csv(data)
    elif lower.endswith(".xls"):
        raise SheetParseError("暂不支持旧版 .xls，请在 Excel 中另存为 .xlsx 后再上传")
    else:
        raise SheetParseError("仅支持 .xlsx 与 .csv 文件")
    total_rows = sum(len(t["rows"]) for t in tables)
    return tables, total_rows
